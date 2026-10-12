import asyncio
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from test_facts_verifier_export import COMPANY, setup_store

from finagent import api, workspace
from finagent.agent import AgentDeps, ResearchAgent
from finagent.config import Settings
from finagent.export.xlsx import quarterly_workbook
from finagent.schemas import ResearchScope
from finagent.tools import generate_financial_workbook


def test_tool_persists_scoped_workbook_and_downloads_verified_sheets(tmp_path, monkeypatch):
    store = setup_store(tmp_path / "tool.sqlite3")
    monkeypatch.setattr(api, "store", lambda: store)
    thread = asyncio.run(workspace.create_thread())
    deps = AgentDeps(
        store=store,
        allowed_citation_ids=set(),
        scope=ResearchScope(companies=[COMPANY], fiscal_years=[2024]),
        tool_context={"thread_id": thread["id"]},
        generated_artifacts=[],
    )
    result = generate_financial_workbook(deps)
    assert result["periods"] == ["2024-Q1", "2024-Q2", "2024-Q3", "2024-Q4"]
    artifact = workspace.thread_data(thread["id"])["artifacts"][0]
    assert artifact["data"]["workbook_layout"] == "quarterly"

    async def download():
        response = await workspace.export_artifact(result["artifact_id"])
        return b"".join([chunk async for chunk in response.body_iterator])

    book = load_workbook(BytesIO(asyncio.run(download())))
    assert set(result["sheets"]) == set(book.sheetnames)
    income = book["季度损益"]
    revenue_column = next(c.column for c in income[1] if c.value == "revenue")
    assert income.cell(5, revenue_column).value == 180
    assert income.cell(5, 2).value == "2024-Q4"
    assert book["Provenance"].max_row > 1
    assert book["数据与计算口径"].max_row > 1
    generate_financial_workbook(deps)
    assert len(deps.generated_artifacts) == 1

    from finagent.schemas import ComparisonTable

    table = ComparisonTable.model_validate(artifact["data"])
    table.verified = False
    with pytest.raises(ValueError):
        quarterly_workbook(table, store)


def test_no_available_periods_does_not_create_artifact(tmp_path, monkeypatch):
    store = setup_store(tmp_path / "missing.sqlite3")
    monkeypatch.setattr(api, "store", lambda: store)
    thread = asyncio.run(workspace.create_thread())
    deps = SimpleNamespace(
        store=store,
        scope=ResearchScope(companies=[COMPANY], fiscal_years=[2030]),
        tool_context={"thread_id": thread["id"]},
        generated_artifacts=[],
    )
    with pytest.raises(ValueError, match="No verified"):
        generate_financial_workbook(deps)
    assert workspace.thread_data(thread["id"])["artifacts"] == []


def test_model_calls_registered_workbook_tool(tmp_path, monkeypatch):
    from pydantic_ai.models.function import DeltaToolCall, FunctionModel

    from finagent.data.facts import FactResolver
    from finagent.schemas import Answer, Claim, SearchHit

    store = setup_store(tmp_path / "model-tool.sqlite3")
    monkeypatch.setattr(api, "store", lambda: store)
    thread = asyncio.run(workspace.create_thread())
    citation = FactResolver(store).resolve(COMPANY, "revenue", "2024-Q1").citations[0]
    calls = 0

    async def model_stream(messages, info):
        nonlocal calls
        calls += 1
        if calls == 1:
            assert {"generate_workbook", "list_files", "read_file", "write_file"}.issubset(
                {tool.name for tool in info.function_tools}
            )
            yield {
                0: DeltaToolCall(
                    name="generate_workbook",
                    json_args='{"frequency":"quarterly"}',
                )
            }
        else:
            answer = Answer(
                claims=[
                    Claim(
                        text="季度工作簿已生成，可在产物区下载。",
                        citation_ids=[citation.citation_id],
                    )
                ]
            )
            yield {
                0: DeltaToolCall(
                    name=info.output_tools[0].name,
                    json_args=answer.model_dump_json(),
                )
            }

    monkeypatch.setattr(
        "pydantic_ai.models.openai.OpenAIResponsesModel",
        lambda *args, **kwargs: FunctionModel(stream_function=model_stream),
    )
    agent = ResearchAgent(Settings(llm_api_key="fixture", llm_model="fixture"), store)
    agent.tool_context = {"thread_id": thread["id"]}

    async def run():
        return [
            item
            async for item in agent.stream_answer(
                "生成季度xlsx",
                ResearchScope(companies=[COMPANY], fiscal_years=[2024]),
                [SearchHit(citation=citation, text="Verified revenue", score=1)],
            )
        ]

    snapshots = asyncio.run(run())
    assert snapshots[-1].status == "complete"
    assert calls == 2
    assert len(agent.generated_artifacts) == 1
    assert all(item.status_reason != "workbook_running" for item in snapshots)
