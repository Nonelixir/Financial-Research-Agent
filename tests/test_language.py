import asyncio
import re
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from test_facts_verifier_export import COMPANY, setup_store
from test_streaming import evidence

from finagent import api, workspace
from finagent.agent import ResearchAgent
from finagent.config import Settings
from finagent.data.facts import FactResolver
from finagent.localization import status_payload
from finagent.prompts import load_prompts
from finagent.schemas import Answer, Claim, ResearchScope
from finagent.tools import generate_financial_workbook, publish_research_report
from finagent.verify.citations import CitationVerificationError, verify_answer_citations


def test_english_prompt_overrides_question_language():
    prompts = load_prompts()
    prompt = prompts.system_for("en")
    assert "Regardless of the language" in prompt
    assert not re.search(r"[\u3400-\u9fff]", prompt)
    assert "中文模式" in prompts.system_for("zh")
    assert ResearchScope().language == "zh"
    with pytest.raises(ValueError):
        ResearchScope(language="fr")


def test_english_fallback_and_event_preserve_diagnostics(tmp_path):
    store, scope, _ = evidence(tmp_path)
    scope.language = "en"
    agent = ResearchAgent(Settings(llm_api_key="", llm_model=""), store)
    answer = asyncio.run(agent.answer("请分析财报", scope, []))
    assert "Insufficient evidence" in answer.caveats[0]
    event = {"phase": "read_file", "state": "warning", "message": "文件操作失败"}
    localized = status_payload("status", event, "en")
    assert localized["state"] == "warning"
    assert localized["message"] == "read_file: failed or limited"
    assert localized["detail"]["original_message"] == event["message"]
    assert status_payload("answer.snapshot", event, "en") is event
    assert status_payload("status", event, "zh") is event


def test_english_workbook_report_and_scope_persistence(tmp_path, monkeypatch):
    store = setup_store(tmp_path / "english.sqlite3")
    monkeypatch.setattr(api, "store", lambda: store)
    thread = asyncio.run(workspace.create_thread())
    scope = ResearchScope(companies=[COMPANY], fiscal_years=[2024], language="en")
    citation = FactResolver(store).resolve(COMPANY, "revenue", "2024-FY").citations[0]
    deps = SimpleNamespace(
        store=store,
        scope=scope,
        allowed_citation_ids={citation.citation_id},
        generated_artifacts=[],
        tool_context={"thread_id": thread["id"]},
    )
    result = generate_financial_workbook(deps)
    artifact = workspace.thread_data(thread["id"])["artifacts"][0]
    assert artifact["scope"]["language"] == "en"
    assert result["title"].startswith("Quarterly financial")
    book = load_workbook(
        BytesIO(workspace.session_files(thread["id"]).read(artifact["file"]["relative_path"]))
    )
    assert set(book.sheetnames) == set(result["sheets"])
    assert book["Quarterly income"].cell(1, 1).value == "Company"
    assert book["Provenance"].max_row > 1
    report = publish_research_report(
        deps,
        "Financial report",
        [Claim(text="Revenue was USD 600.", citation_ids=[citation.citation_id])],
        ["Limited sample."],
    )
    saved = workspace.thread_data(thread["id"])["artifacts"][-1]
    assert saved["id"] == report["artifact_id"]
    assert "## Sources" in saved["data"]["text"]
    assert "## Limitations and methodology" in saved["data"]["text"]


@pytest.mark.parametrize("amount", ["USD 600", "$600", "0.0006 million USD", "US$0.6 thousand"])
def test_english_currency_scales(tmp_path, amount):
    store = setup_store(tmp_path / "usd.sqlite3")
    citation = FactResolver(store).resolve(COMPANY, "revenue", "2024-FY").citations[0]
    ids = {citation.citation_id}
    verify_answer_citations(
        Answer(claims=[Claim(text=f"Revenue was {amount}.", citation_ids=list(ids))]), ids, store
    )
    with pytest.raises(CitationVerificationError, match="unit mismatch"):
        verify_answer_citations(
            Answer(claims=[Claim(text="Revenue was USD 600 billion.", citation_ids=list(ids))]),
            ids,
            store,
        )
