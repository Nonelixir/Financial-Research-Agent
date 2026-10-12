from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from finagent.schemas import ComparisonTable, VerificationStatus
from finagent.storage import SQLiteStore


def comparison_workbook(table: ComparisonTable, store: SQLiteStore) -> bytes:
    if not table.verified or any(
        row.cell.verification.status == VerificationStatus.MISMATCH for row in table.rows
    ):
        raise ValueError("Verification failed; workbook export is blocked")
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    provenance = workbook.create_sheet("Provenance")
    data.append(
        ["Company", "Ticker", "Metric", "Period", "Value", "Unit", "Status", "Citation IDs"]
    )
    provenance.append(
        [
            "Citation ID",
            "Ticker",
            "Accession",
            "Form",
            "Filed",
            "Locator",
            "Value",
            "Unit",
            "SEC URL",
        ]
    )
    header_fill = PatternFill("solid", fgColor="0F766E")
    for sheet in (data, provenance):
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
    exported: set[str] = set()
    for row in table.rows:
        cell = row.cell
        data.append(
            [
                row.company.name,
                row.company.ticker,
                row.metric,
                row.period,
                cell.value,
                cell.unit,
                cell.verification.status.value,
                ", ".join(cell.source_citation_ids),
            ]
        )
        for citation_id in cell.source_citation_ids:
            if citation_id in exported:
                continue
            citation = store.get_citation(citation_id)
            if citation is None:
                raise ValueError(f"Missing citation {citation_id}")
            exported.add(citation_id)
            provenance.append(
                [
                    citation.citation_id,
                    citation.source.ticker,
                    citation.source.accession_no,
                    citation.source.form,
                    citation.source.filing_date.isoformat(),
                    citation.locator.model_dump_json(),
                    citation.value.amount if citation.value else None,
                    citation.value.unit if citation.value else None,
                    str(citation.source.sec_url),
                ]
            )
            url_cell = provenance.cell(provenance.max_row, 9)
            url_cell.hyperlink = str(citation.source.sec_url)
            url_cell.style = "Hyperlink"
    for sheet in (data, provenance):
        for column in sheet.columns:
            width = min(60, max(len(str(cell.value or "")) for cell in column) + 2)
            sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def quarterly_workbook(table: ComparisonTable, store: SQLiteStore, language: str = "zh") -> bytes:
    """Group the verified snapshot by statement, retaining full provenance."""
    from openpyxl import load_workbook

    from finagent.data.facts import INSTANT_METRICS

    workbook = load_workbook(BytesIO(comparison_workbook(table, store)))
    workbook.remove(workbook["Data"])
    cash_metrics = {
        "operating_cash_flow",
        "investing_cash_flow",
        "financing_cash_flow",
        "capital_expenditure",
    }
    groups = {
        "季度损益": [r for r in table.rows if r.metric not in INSTANT_METRICS | cash_metrics],
        "季度现金流": [r for r in table.rows if r.metric in cash_metrics],
        "季度资产负债表": [r for r in table.rows if r.metric in INSTANT_METRICS],
    }
    for title, rows in groups.items():
        sheet = workbook.create_sheet(title)
        metrics = list(dict.fromkeys(r.metric for r in rows))
        sheet.append(["公司", "季度", "单位", *metrics])
        periods = sorted({(r.company.ticker, r.period) for r in rows})
        lookup = {(r.company.ticker, r.period, r.metric): r.cell for r in rows}
        for ticker, period in periods:
            sheet.append(
                [
                    ticker,
                    period,
                    "USD",
                    *[
                        lookup[(ticker, period, metric)].value
                        if (ticker, period, metric) in lookup
                        else None
                        for metric in metrics
                    ],
                ]
            )
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="0F766E")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = max(
                16, min(40, len(str(column[0].value)) + 2)
            )
        for row in sheet.iter_rows(min_row=2, min_col=4):
            for cell in row:
                cell.number_format = "#,##0.00;[Red](#,##0.00)"
    lineage = workbook.create_sheet("数据与计算口径")
    lineage.append(["公司", "指标", "季度", "值", "单位", "核验状态", "公式", "引用 ID"])
    for row in table.rows:
        cell = row.cell
        lineage.append(
            [
                row.company.ticker,
                row.metric,
                row.period,
                cell.value,
                cell.unit,
                cell.verification.status.value,
                cell.formula.expression if cell.formula else "",
                ", ".join(cell.source_citation_ids),
            ]
        )
    lineage.freeze_panes = "A2"
    notes = workbook.create_sheet("说明")
    notes.append(["口径", "说明"])
    notes.append(["金额", "USD 原币单位；空白表示缺失，不是零。"])
    notes.append(["期间", "损益和现金流为单季度流量；资产负债表为期末余额。"])
    notes.append(["派生", "Q4=全年减Q3累计；累计现金流差分见数据与计算口径。"])
    for warning in table.warnings:
        notes.append(["数据提示", warning])
    if language == "en":
        labels = {
            "季度损益": "Quarterly income",
            "季度现金流": "Quarterly cash flow",
            "季度资产负债表": "Quarterly balance sheet",
            "数据与计算口径": "Data and calculations",
            "说明": "Notes",
            "公司": "Company",
            "季度": "Quarter",
            "单位": "Unit",
            "指标": "Metric",
            "值": "Value",
            "核验状态": "Verification",
            "公式": "Formula",
            "引用 ID": "Citation IDs",
            "口径": "Methodology",
            "金额": "Amounts",
            "期间": "Periods",
            "派生": "Derivation",
            "数据提示": "Data warning",
            "USD 原币单位；空白表示缺失，不是零。": (
                "Original USD units; blank means missing, not zero."
            ),
            "损益和现金流为单季度流量；资产负债表为期末余额。": (
                "Income and cash flow are standalone quarters; "
                "balance sheets are period-end snapshots."
            ),
            "Q4=全年减Q3累计；累计现金流差分见数据与计算口径。": (
                "Q4 = FY minus Q3 YTD; see Data and calculations for cash-flow differencing."
            ),
        }
        for sheet in workbook:
            sheet.title = labels.get(sheet.title, sheet.title)
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = labels.get(cell.value, cell.value)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 100
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
